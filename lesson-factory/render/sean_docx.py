"""LessonPlan JSON → 세안 `.docx`.

조판 데이터는 sean.py 가 만든다. 이 파일은 그것을 워드 표에 앉히는 일만 한다.

CLAUDE.md B-2: 파일 쓰기는 dry_run 이 기본이다. 실제로 저장하려면 --write 를 준다.

    python3 -m render.sean_docx tests/fixtures/사회-5-2-2단원-1차시.plan.json
    python3 -m render.sean_docx tests/fixtures/사회-5-2-2단원-1차시.plan.json --write -o artifacts/
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from core import paths

from . import sean

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_DB = ROOT / "curriculum" / "standards.sqlite"

# 본시 과정 표 5열의 너비 비율. 실물 세안은 활동 칸이 압도적으로 넓다.
PROCESS_WIDTHS = (0.10, 0.15, 0.52, 0.06, 0.17)


def _set_cell(cell, text: str, *, bold: bool = False, size: int = 9) -> None:
    from docx.shared import Pt

    cell.text = ""
    para = cell.paragraphs[0]
    for i, line in enumerate(text.split("\n")):
        if i:
            para = cell.add_paragraph()
        run = para.add_run(line)
        run.bold = bold
        run.font.size = Pt(size)
        run.font.name = "맑은 고딕"


def _table(doc, rows: int, cols: int):
    table = doc.add_table(rows=rows, cols=cols)
    table.style = "Table Grid"
    return table


def build(plan: dict, db_path: Path | None = None):
    """python-docx Document 를 만들어 돌려준다. 파일로 쓰지 않는다."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    meta = plan["meta"]
    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(0.6)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(f"{meta['subject']}과 교수·학습 과정안 (세안)")
    run.bold = True
    run.font.size = Pt(16)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = subtitle.add_run(
        f"{meta['grade']}학년 {meta['semester']}학기 · {meta['unit'].split('.')[0].strip()}단원 "
        f"{(meta.get('sub_unit') or '').split()[0]} · {meta['period'].split('/')[0]}차시"
        + (f" | {plan['teaching_model']}" if plan.get("teaching_model") else "")
    )
    sub.font.size = Pt(10)

    # 머리표 — (라벨, 값) 두 쌍씩
    pairs = sean.header_pairs(plan, db_path)
    head = _table(doc, len(pairs), 4)
    for row, (la, va, lb, vb) in zip(head.rows, pairs):
        _set_cell(row.cells[0], la, bold=True)
        _set_cell(row.cells[1], va)
        _set_cell(row.cells[2], lb, bold=True)
        _set_cell(row.cells[3], vb)

    doc.add_paragraph()
    doc.add_paragraph().add_run("■ 본시 교수·학습 과정").bold = True

    rows = sean.process_rows(plan)
    process = _table(doc, len(rows) + 1, 5)
    headers = ("단계(분)", "학습 내용", "교수·학습 활동 (T 교사 / S 학생 예상 반응)", "시간", "자료(★)·유의점(※)")
    for cell, text in zip(process.rows[0].cells, headers):
        _set_cell(cell, text, bold=True)
    for row, data in zip(process.rows[1:], rows):
        for cell, text in zip(row.cells, data):
            _set_cell(cell, text)

    total = Inches(9.0)
    for row in process.rows:
        for cell, ratio in zip(row.cells, PROCESS_WIDTHS):
            cell.width = int(total * ratio)

    assessment = sean.assessment_rows(plan)
    if assessment:
        doc.add_paragraph()
        doc.add_paragraph().add_run("■ 평가 계획").bold = True
        table = _table(doc, len(assessment) + 1, 4)
        for cell, text in zip(table.rows[0].cells, ("평가 요소", "평가 기준", "방법", "시기")):
            _set_cell(cell, text, bold=True)
        for row, data in zip(table.rows[1:], assessment):
            for cell, text in zip(row.cells, data):
                _set_cell(cell, text)

    doc.add_paragraph()
    doc.add_paragraph().add_run("■ 판서 계획").bold = True
    for line in sean.board_plan(plan):
        doc.add_paragraph(line)

    if plan.get("guidance_notes"):
        doc.add_paragraph()
        doc.add_paragraph().add_run("■ 지도상의 유의점").bold = True
        for i, note in enumerate(plan["guidance_notes"]):
            doc.add_paragraph(f"{sean._circled(i)} {note}")

    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    mark = footer.add_run(plan["copyright_notice"])
    mark.font.size = Pt(8)

    return doc


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="LessonPlan JSON → 세안 .docx")
    ap.add_argument("plan", type=Path)
    ap.add_argument("-o", "--out", type=Path, default=None, help="산출물 뿌리 (기본: .env 의 LESSON_OUT 또는 artifacts/)")
    ap.add_argument("--db", type=Path, default=DEFAULT_DB)
    ap.add_argument("--write", action="store_true", help="실제로 파일을 쓴다 (기본은 dry_run)")
    args = ap.parse_args(argv)

    plan = json.loads(args.plan.read_text(encoding="utf-8"))
    target = paths.lesson_dir(plan["trace_id"], args.out, subject=plan["meta"]["subject"]) / f"{plan['trace_id']}-세안-v1.docx"

    if not args.write:
        print(sean.render_markdown(plan, args.db if args.db.exists() else None))
        print(f"\n[dry_run] 저장하려면 --write. 저장 위치: {target}", file=sys.stderr)
        return 0

    doc = build(plan, args.db if args.db.exists() else None)
    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(target)
    print(f"저장함: {target}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
